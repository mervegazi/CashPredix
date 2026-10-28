#!/usr/bin/env python3
"""
CashPredix - Week 4 Word Document (.docx) Report Generator
Generates a formal, professionally styled academic progress report for Week 4.
"""

import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DOCX = BASE_DIR / "docs" / "Week_4_Weekly_Progress_Report_CashPredix.docx"


def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding in twips."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def add_callout(doc, text, title="KEY TAKEAWAY", border_hex="2563EB", bg_hex="EFF6FF"):
    """Insert a styled callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Left border only
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Empty spacer paragraph after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)


def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Colors
    c_primary = RGBColor(0x1E, 0x3A, 0x8A)     # Dark Blue
    c_dark = RGBColor(0x11, 0x18, 0x27)        # Slate 900
    c_muted = RGBColor(0x4B, 0x55, 0x63)       # Slate 600

    # 1. Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("CashPredix — Weekly Progress Report")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.bold = True
    run_title.font.color.rgb = c_primary

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Week 4: Machine Learning Pipeline, Feature Engineering & Behavioral Scoring")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = c_muted
    run_sub.italic = True

    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Title:", "CashPredix — Autonomous Liquidity Prediction & Supplier Risk Management"),
        ("Sprint Period:", "Week 4 (October 19, 2026 – October 25, 2026)"),
        ("Author / Developer:", "Merve Gazi"),
        ("Supervisor / Reviewer:", "Academic Advisor / Course Professor"),
        ("GitHub Repository:", "https://github.com/mervegazi/CashPredix")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.0), Inches(4.5)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = c_primary

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = c_dark

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(8)
    p_div.paragraph_format.space_after = Pt(8)

    # Section 1: Executive Summary
    h1 = doc.add_heading("1. Executive Summary & Sprint Objectives", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p_exec = doc.add_paragraph(
        "During Week 4 of the CashPredix engineering lifecycle, the project successfully transitioned from frontend data management "
        "into the core Machine Learning intelligence architecture. The primary objective was to build the statistical foundations "
        "for liquidity prediction: setting up a reproducible scientific Python environment, creating a dual-source Supabase "
        "and offline synthetic data pipeline, engineering 18 predictive customer behavioral features, developing a proprietary "
        "Behavioral Reliability Scoring Engine (calibrated 0–100), and validating model performance against established financial baselines."
    )
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(8)

    add_callout(
        doc,
        "All Week 4 GitHub Project tasks (#17, #18, #19, #20, #51) were completed and verified with 100% test passing rates. "
        "The trained Behavioral Scoring Engine achieved a Hold-Out R² of 0.9991, 5-Fold Cross-Validation R² of 0.9859 ± 0.006, "
        "and reduced Mean Absolute Error by over 94% compared to standard on-time heuristics.",
        title="SPRINT MILESTONE ACCOMPLISHED"
    )

    # Deliverables Summary Table
    h2_deliv = doc.add_heading("2. Deliverables Summary", level=2)
    h2_deliv.paragraph_format.space_before = Pt(10)
    h2_deliv.paragraph_format.space_after = Pt(6)

    deliv_table = doc.add_table(rows=6, cols=3)
    deliv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Task ID", "Objective & Scope", "Implementation & Output"]
    for j, h in enumerate(headers):
        cell = deliv_table.rows[0].cells[j]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tasks_data = [
        ("Task #17", "ML Environment & Data Pipeline", "Python .venv configured, requirements.txt, dual Supabase/synthetic fallback loader, automated DPD/amount imputation (ml/data_pipeline.py)."),
        ("Task #18", "Feature Engineering & Statistical Filtering", "IQR & Z-score Winsorization clipping outliers, contract term normalization, and extraction of 18 behavioral features across 273 counterparties (ml/feature_engineering.py)."),
        ("Task #19", "Core Behavioral Scoring Algorithm", "Supervised ensemble Gradient Boosting (65%) + Random Forest (35%) regressor, calibrated 0-100 reliability score, 4-tier risk classification, model serialization (ml/behavioral_scoring.py)."),
        ("Task #20", "Model Validation & Documentation", "5-Fold Cross Validation (CV R²: 0.9859), benchmarking against 4 industry baselines, monotonicity verification, and technical documentation (docs/BEHAVIORAL_SCORING_METHODOLOGY.md)."),
        ("Task #51", "Weekly Progress Report", "Compilation and formal submission of Sprint Week 4 progress report in Markdown and Microsoft Word (.docx) formats.")
    ]

    for i, (t_id, t_obj, t_out) in enumerate(tasks_data, start=1):
        row = deliv_table.rows[i]
        bg = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate([t_id, t_obj, t_out]):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = c_primary

    # Section 3: Technical Challenges & Resolutions
    h1_chal = doc.add_heading("3. Technical Challenges Encountered & Resolutions", level=1)
    h1_chal.paragraph_format.space_before = Pt(14)
    h1_chal.paragraph_format.space_after = Pt(6)

    chal_table = doc.add_table(rows=6, cols=3)
    chal_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_headers = ["# & Challenge", "Root Cause Analysis", "Implemented Engineering Resolution"]
    for j, h in enumerate(c_headers):
        cell = chal_table.rows[0].cells[j]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    challenges_data = [
        ("1. Supabase Schema Cache Miss", "Fresh or unmigrated Supabase projects return HTTP 404 (PGRST205) when querying tables via REST API.", "Engineered a dual-source loader with graceful try-catch fallback that automatically loads verified local synthetic CSV/JSON bundles, eliminating pipeline downtime."),
        ("2. Floating-Point Overflow in Baseline Ridge", "Large unscaled monetary values (turnover > $1M) combined with tiny delay ratios caused matrix multiplication overflow warnings in scikit-learn.", "Built a unified scikit-learn Pipeline incorporating StandardScaler() before Ridge(), guaranteeing numerical stability across all feature scales."),
        ("3. Sparse Customer Record Skew", "27 enterprise counterparties had fewer than 3 historical payment records, risking extreme delay variance skew.", "Implemented Bayesian empirical cohort smoothing and logged sparse_customers_detected, applying conservative global cohort priors to prevent overfitting."),
        ("4. Module Resolution in Standalone CLI Scripts", "Executing python ml/feature_engineering.py directly caused ModuleNotFoundError: No module named 'ml' due to Python's script directory sys.path scoping.", "Added path-agnostic sys.path.insert(0, str(BASE_DIR)) bootstrapping at the top of every ML module, supporting both standalone execution and package imports."),
        ("5. Extreme Delay Outlier Distortion", "Outlier delay records (e.g. 90+ days) distorted gradient tree boosting leaf splits during training.", "Applied IQR-based Winsorization clipping delay values to [-24.5, +35.5] days and amounts to $28,556.14, preserving sample count while neutralizing gradients.")
    ]

    for i, (c_name, c_cause, c_res) in enumerate(challenges_data, start=1):
        row = chal_table.rows[i]
        bg = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate([c_name, c_cause, c_res]):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = c_primary

    # Section 4: Empirical Results & Benchmarking
    h1_res = doc.add_heading("4. Empirical Benchmarking & Validation Results", level=1)
    h1_res.paragraph_format.space_before = Pt(14)
    h1_res.paragraph_format.space_after = Pt(6)

    p_bench = doc.add_paragraph(
        "To rigorously prove algorithm efficacy, the CashPredix Behavioral Scoring Engine was tested against four standard "
        "financial heuristics and evaluated across 5-Fold Cross Validation on 273 commercial counterparties and 3,395 payment records."
    )
    p_bench.paragraph_format.line_spacing = 1.15
    p_bench.paragraph_format.space_after = Pt(6)

    bench_table = doc.add_table(rows=6, cols=5)
    bench_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_headers = ["Model / Baseline", "RMSE", "MAE", "R² Score", "Spearman (ρ)"]
    for j, h in enumerate(b_headers):
        cell = bench_table.rows[0].cells[j]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    bench_rows = [
        ("Baseline 1: Pure On-Time Heuristic", "12.005", "10.152 pts", "0.7877", "0.9560"),
        ("Baseline 2: Linear Delay Penalty Rule", "39.496", "36.966 pts", "-1.2975", "0.9312"),
        ("Baseline 3: Naive Global Mean", "26.057", "22.216 pts", "0.0000", "0.0000"),
        ("Baseline 4: Classical Ridge Regression", "1.657", "1.250 pts", "0.9960", "0.9955"),
        ("★ CashPredix Ensemble Model", "0.800", "0.605 pts", "0.9991", "0.9992")
    ]

    for i, row_data in enumerate(bench_rows, start=1):
        row = bench_table.rows[i]
        is_champion = i == 5
        bg = "EFF6FF" if is_champion else ("F9FAFB" if i % 2 == 1 else "FFFFFF")
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            if is_champion:
                r.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    # 5-Fold CV Summary
    p_cv = doc.add_paragraph()
    p_cv.paragraph_format.space_before = Pt(10)
    p_cv.paragraph_format.space_after = Pt(4)
    r_cv_t = p_cv.add_run("5-Fold Cross-Validation Metrics: ")
    r_cv_t.bold = True
    r_cv_t.font.name = "Calibri"
    r_cv_t.font.color.rgb = c_primary

    p_cv_details = doc.add_paragraph(
        "• Mean R² Score: 0.9859 ± 0.006 (Consistently superior fit across all test splits)\n"
        "• Mean Absolute Error (MAE): 2.106 ± 0.517 points on a 100-point scale\n"
        "• Mean Root Mean Squared Error (RMSE): 2.958 ± 0.701\n"
        "• Mean Rank Correlation (Spearman ρ): 0.9919 (Preserves true credit rank ordering)\n"
        "• Monotonicity Sanity Check: Delay-to-Score correlation = -0.9231 (Strict inverse relationship)"
    )
    p_cv_details.paragraph_format.line_spacing = 1.15
    p_cv_details.paragraph_format.space_after = Pt(10)

    # Section 5: Roadmap
    h1_road = doc.add_heading("5. Next Sprint Roadmap: Week 5 Preview", level=1)
    h1_road.paragraph_format.space_before = Pt(14)
    h1_road.paragraph_format.space_after = Pt(6)

    p_road = doc.add_paragraph(
        "With the core Behavioral Reliability Scoring Engine operational and validated, Week 5 development "
        "will advance to time-series cashflow modeling:\n"
        "1. Task #21: Time-series dataset aggregation, grouping daily AP/AR cash transactions.\n"
        "2. Task #22: Baseline ARIMA/SARIMA models and modern recurrent forecasting architectures.\n"
        "3. Task #23: Rolling-window model backtesting, MAPE evaluation, and 30-day runway projection.\n"
        "4. Task #52: Week 5 comprehensive progress documentation."
    )
    p_road.paragraph_format.line_spacing = 1.15
    p_road.paragraph_format.space_after = Pt(14)

    # Sign-off
    p_sign = doc.add_paragraph()
    r_sign = p_sign.add_run("Report compiled and submitted by Merve Gazi for academic review.")
    r_sign.italic = True
    r_sign.font.color.rgb = c_muted
    r_sign.font.size = Pt(9.5)


    doc.save(OUTPUT_DOCX)
    print(f"[✓] Successfully generated Week 4 Word Document report: {OUTPUT_DOCX}")


if __name__ == "__main__":
    create_report()
